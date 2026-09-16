from pathlib import Path

import docx
from langchain_core.documents import Document
from pypdf import PdfReader
from langchain_text_splitters import RecursiveCharacterTextSplitter

from core.config import settings


SUPPORTED_SUFFIXES = {".pdf", ".docx", ".md", ".txt"}
VISIBILITIES = {"public", "internal"}


def _visibility(path: Path, docs_dir: Path) -> str:
    relative = path.relative_to(docs_dir)
    first_part = relative.parts[0].lower() if len(relative.parts) > 1 else "public"
    if first_part in {"employee", "hr", "internal"}:
        return "internal"
    return first_part if first_part in VISIBILITIES else "public"


def _metadata(path: Path, docs_dir: Path, **extra) -> dict:
    return {
        "source": path.name,
        "source_path": path.relative_to(docs_dir).as_posix(),
        "visibility": _visibility(path, docs_dir),
        "file_type": path.suffix.lower().lstrip("."),
        **extra,
    }


def load_pdf(path: Path, docs_dir: Path) -> list[Document]:
    reader = PdfReader(str(path))
    documents: list[Document] = []
    for page_number, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        if text.strip():
            documents.append(
                Document(
                    page_content=text,
                    metadata=_metadata(path, docs_dir, page=page_number),
                )
            )
    return documents


def load_docx(path: Path, docs_dir: Path) -> list[Document]:
    document = docx.Document(str(path))
    blocks = [paragraph.text.strip() for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            blocks.append(" | ".join(cell.text.strip() for cell in row.cells))
    text = "\n".join(block for block in blocks if block)
    if not text:
        return []
    return [Document(page_content=text, metadata=_metadata(path, docs_dir))]


def load_text(path: Path, docs_dir: Path) -> list[Document]:
    text = path.read_text(encoding="utf-8-sig")
    if not text.strip():
        return []
    return [Document(page_content=text, metadata=_metadata(path, docs_dir))]


def load_file(path: Path, docs_dir: Path) -> list[Document]:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return load_pdf(path, docs_dir)
    if suffix == ".docx":
        return load_docx(path, docs_dir)
    if suffix in {".txt", ".md"}:
        return load_text(path, docs_dir)
    raise ValueError("只支持 txt、md、pdf、docx")


def load_docs(dir_path: str | Path | None = None) -> list[Document]:
    docs_dir = Path(dir_path or settings.docs_dir).resolve()
    if not docs_dir.is_dir():
        raise FileNotFoundError(f"知识文档目录不存在：{docs_dir}")

    legacy_docs = sorted(docs_dir.rglob("*.doc"))
    if legacy_docs:
        names = "、".join(path.name for path in legacy_docs)
        raise ValueError(f"不支持旧版 .doc，请先转换为 .docx：{names}")

    documents: list[Document] = []
    for path in sorted(docs_dir.rglob("*")):
        if not path.is_file() or path.suffix.lower() not in SUPPORTED_SUFFIXES:
            continue
        documents.extend(load_file(path, docs_dir))
    return documents


def split_docs(documents: list[Document]) -> list[Document]:
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=settings.chunk_size,
        chunk_overlap=settings.chunk_overlap,
        add_start_index=True,
    )
    return splitter.split_documents(documents)